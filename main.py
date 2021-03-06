from tkinter import *
from tkinter import filedialog
from tkinter import messagebox
import codecs
from docx import Document  # 导入库
from datetime import date
import glob
import os
import win32com.client as wc

base = Tk()
base.title("api接口解析生成java类文件工具")
base.geometry('900x900')


# Function for opening the file
def get_file_path():
    file_path_name.set(filedialog.askopenfile().name)


def get_target_path():
    target_path_name.set(filedialog.askdirectory())


def doc2docx(file):
    word = wc.Dispatch("Word.Application")
    doc = word.Documents.Open(file)
    # 上面的地方只能使用完整绝对地址，相对地址找不到文件，且，只能用“\\”，不能用“/”，哪怕加了 r 也不行，涉及到将反斜杠看成转义字符。
    doc.SaveAs(file + "x", 12, False, "", True, "", False, False, False,
               False)  # 转换后的文件,12代表转换后为docx文件
    # 注意SaveAs会打开保存后的文件，有时可能看不到，但后台一定是打开的
    doc.Close
    word.Quit


def pdf_to_docx(pdfs_path):
    word = wc.Dispatch("Word.Application")
    word.visible = 0
    for i, doc in enumerate(glob.iglob(pdfs_path)):
        filename = doc.split('\\')[-1]
        filepath = doc.replace(filename, "")
        in_file = os.path.abspath(doc)
        wb = word.Documents.Open(in_file)
        out_file = os.path.abspath(filepath + filename[0:-4] + ".docx".format(i))
        print("outfile\n", out_file)
        wb.SaveAs2(out_file, FileFormat=16)  # file format for docx
        print("pdf to docx success...")
        wb.Close()
    word.Quit()


def start_generate_file(file_and_target_path, param_sort, channel_name_and_coder_name, api_name, flag_name, is_snake_case):
    print(is_snake_case)
    description = [""]
    # 如果是doc文件转换成docx文件后在进行处理
    if str(file_and_target_path[0]).endswith(".doc"):
        doc2docx(file_and_target_path[0])
        file_and_target_path[0] = file_and_target_path[0] + "x"
    # 如果是pdf文转换成docx文件后在进行处理件
    if str(file_and_target_path[0]).endswith(".pdf"):
        pdf_to_docx(file_and_target_path[0])
        file_and_target_path[0] = str(file_and_target_path[0]).replace("pdf", "docx")

    # 进入doc解析核心处理
    parse_docx_file(file_and_target_path[0], channel_name_and_coder_name[0], channel_name_and_coder_name[1],
                    description,
                    api_name, file_and_target_path[1], flag_name, param_sort, is_snake_case)
    messagebox.showinfo('提示', '生成成功，请到指定文件夹下查看')
    pass


# 判断是否有下划线, 有的话转为驼峰格式
def to_camel_case(snake_str, is_snake_case):
    if str(is_snake_case) == "1":
        if "_" in snake_str:
            components = snake_str.split('_')
            # We capitalize the first letter of each component except the first one
            # with the 'title' method and join them together.
            return components[0] + ''.join(x.title() for x in components[1:])
    return snake_str


# 解析docx文件
def parse_docx_file(docx_file_path, channel_name, coder_name, description, class_name, target_directoy, flag_name,
                    param_sort, is_snake_case):
    print("parse_docx_file :" + docx_file_path)
    document = Document(docx_file_path)  # 读入文件
    tables = document.tables  # 获取文件中的表格集
    index = 0
    for table in tables:
        # 找到请求响应类的表格位置, 如果不是就跳出本次循环
        if table.cell(0, 0).text == flag_name:
            # 生成class内容
            content, end, head = generate_class_content(coder_name, index, description, class_name, channel_name)
            for i in range(1, len(table.rows)):  # 从表格第二行开始循环读取表格数据
                annotation = "\t/** " + table.cell(i, int(param_sort[1]) - 1).text + ", " + table.cell(i,
                                                                                                       int(param_sort[
                                                                                                               2]) - 1).text + " **/ \n"
                var = "\tprivate String " + to_camel_case(table.cell(i, int(param_sort[0]) - 1).text, is_snake_case) + ";\n\n"
                content += annotation + var
            all_content = head + content + end
            print(all_content)
            # 将文件生成到指定路径下
            output_java_file(all_content, class_name, index, target_directoy, channel_name)
            index = index + 1


# 将生成的文件输出到指定路径下
def output_java_file(all_content, class_name, index, target_directoy, channel_name):
    file_name = target_directoy + "/" + channel_name + class_name[index] + get_suffix(index) + '.java'
    file = codecs.open(file_name, "w", "utf-8")
    file.write(all_content)


# 生成类的内容
def generate_class_content(coder_name, index, description, class_name, channel_name):
    head = "/** \n* @author " + coder_name + " \n* @description " + "\n* @date " + get_current_time() + \
           "\n*/\n@Data\npublic class " + channel_name + class_name[index] + get_suffix(index) + " { \n\n "
    content = ""
    end = "} \n"
    return content, end, head


# Request和Response后缀循环生成
def get_suffix(index):
    suffix = "Request" if index % 2 == 0 else "Response"
    return suffix


# 获取当前时间
def get_current_time():
    return date.today().strftime("%B %d, %Y")


file_path_name = StringVar()
file_path = Button(base, text='1.---请选择doc/docx/pdf接口文档---', command=lambda: get_file_path())
file_path.pack()
e2 = Entry(base, state='readonly', text=file_path_name, width=80)
e2.pack()

L1 = Label(base, text="渠道英文简称,如Wechat")
L1.pack()
channel_name = Entry(base, bd=5)
channel_name.pack()

L2 = Label(base, text="开发人员名称，如Mr.Cloud")
L2.pack()
coder_name = Entry(base, bd=5)
coder_name.pack()

L3 = Label(base, text="文档中每个表格中的第一个字符串标识(用来定位表格)")
L3.pack()
flag_name = Entry(base, bd=5)
flag_name.pack()

V1 = Label(base, text="英文参数顺序(按表的列顺序填入,如:1)")
V1.pack()
en_param = Entry(base, bd=5)
en_param.pack()

V2 = Label(base, text="中文参数(按表的列顺序填入,如:2)")
V2.pack()
ch_param = Entry(base, bd=5)
ch_param.pack()

V3 = Label(base, text="说明(按表格的列顺序填入，如:3)")
V3.pack()
explain = Entry(base, bd=5)
explain.pack()

L3 = Label(base, text="接口类型名称1,如Micropay")
L3.pack()
api_name1 = Entry(base, bd=5)
api_name1.pack()

L4 = Label(base, text="接口类型名称2,如Refund")
L4.pack()
api_name2 = Entry(base, bd=5)
api_name2.pack()

L5 = Label(base, text="接口类型名称3,如OrderQuery")
L5.pack()
api_name3 = Entry(base, bd=5)
api_name3.pack()

L6 = Label(base, text="接口类型名称4,如Reverse")
L6.pack()
api_name4 = Entry(base, bd=5)
api_name4.pack()

L7 = Label(base, text="接口类型名称5,如MerchIn")
L7.pack()
api_name5 = Entry(base, bd=5)
api_name5.pack()

CheckVar1 = IntVar()
C1 = Checkbutton(base, text="是否将下划线转成驼峰", variable=CheckVar1, onvalue=1, offvalue=0, height=5, width=20)
C1.pack()

target_path_name = StringVar()
target_path = Button(base, text='2.---请选择java文件生成的位置---', command=lambda: get_target_path())
target_path.pack()
e1 = Entry(base, state='readonly', text=target_path_name, width=80)
e1.pack()

submit = Button(base, text='--生成--',
                command=lambda: start_generate_file([e2.get(), e1.get()],
                                                    [en_param.get(), ch_param.get(), explain.get()],
                                                    [channel_name.get(), coder_name.get()],
                                                    [api_name1.get(), api_name1.get(), api_name2.get(), api_name2.get(),
                                                     api_name3.get(), api_name3.get(), api_name4.get(), api_name4.get(),
                                                     api_name5.get(), api_name5.get()], flag_name.get(), CheckVar1.get()))

submit.pack()

if __name__ == '__main__':
    mainloop()
